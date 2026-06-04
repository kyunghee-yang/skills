"""적대적 입력 종단 통합 테스트.

수식 인젝션 가맹점 / 빈 시간 / 비숫자 금액이 섞인 승인내역 .xls 로 run_pipeline 전체를
돌려, 누적 하드닝(시간 파싱·금액 파싱·수식 인젝션)이 함께 작동해 크래시 없이 안전한 결의서를
생성함을 검증한다.
"""
import os
import sys

import openpyxl
import xlwt

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "src"))

from expense_report.main import run_pipeline
from expense_report.config import (
    XLS_DATA_START_ROW, XLS_COL_DATE, XLS_COL_TIME, XLS_COL_MERCHANT,
    XLS_COL_CARD_NUMBER, XLS_COL_TYPE, XLS_COL_AMOUNT, XLS_COL_TRANSACTION_TYPE,
    XLS_COL_APPROVAL_NUMBER, XLS_COL_PURCHASE, XLS_COL_PURCHASE_DATE,
    XLS_COL_INSTALLMENT, XLS_COL_VAT_OR_STATUS,
)

# (date, time, merchant, amount) — 적대적 행들
_ADVERSARIAL = [
    ("2026.03.10", "", '=HYPERLINK("http://evil","x")', "9,000"),   # 수식 인젝션 + 빈 시간
    ("2026.03.11", "12:30", "쿠팡", "N/A"),                          # 비숫자 금액
    ("2026.03.12", "비정상", "정상가맹점", "10,000"),                  # 비정상 시간
]


def _build_adversarial_xls(path):
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("승인내역")
    # 메타/헤더 영역 채우기
    for col in range(XLS_COL_VAT_OR_STATUS + 1):
        ws.write(XLS_DATA_START_ROW - 1, col, f"h{col}")
    row = XLS_DATA_START_ROW
    for date, time, merchant, amount in _ADVERSARIAL:
        ws.write(row, XLS_COL_DATE, date)
        ws.write(row, XLS_COL_TIME, time)
        ws.write(row, XLS_COL_MERCHANT, merchant)
        ws.write(row, XLS_COL_CARD_NUMBER, "4201-****-****-7592")
        ws.write(row, XLS_COL_TYPE, "국내일반")
        ws.write(row, XLS_COL_AMOUNT, amount)
        ws.write(row, XLS_COL_TRANSACTION_TYPE, "국내 일시불")
        ws.write(row, XLS_COL_APPROVAL_NUMBER, "1")
        ws.write(row, XLS_COL_PURCHASE, "매입")
        ws.write(row, XLS_COL_PURCHASE_DATE, date)
        ws.write(row, XLS_COL_INSTALLMENT, "-")
        ws.write(row, XLS_COL_VAT_OR_STATUS, "0")
        ws.write(row + 1, XLS_COL_VAT_OR_STATUS, "정상")
        row += 2
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def test_pipeline_survives_adversarial_input(tmp_path, sample_template):
    sub = tmp_path / "202603"
    sub.mkdir()
    _build_adversarial_xls(str(sub / "간편서비스_승인내역.xls"))

    # 크래시 없이 완주해야 한다(누적 하드닝)
    result = run_pipeline(str(sub), notion_data=None, template_path=sample_template)
    assert os.path.exists(result["created_file"])
    assert result["total_count"] == 3

    # 수식 인젝션이 결의서에서 무력화됐는지(가맹점 셀이 수식 아님)
    wb = openpyxl.load_workbook(result["created_file"])
    ws = wb["1.매출내역(원본)"]
    found_evil = False
    for r in range(10, 10 + 3 * 2):
        v = ws.cell(r, 7).value  # SHEET1_COL_MERCHANT
        if v and "HYPERLINK" in str(v):
            found_evil = True
            assert ws.cell(r, 7).data_type != "f"      # 수식 아님
            assert str(v).startswith("'=")             # 텍스트로 살균
    assert found_evil  # 악성 가맹점이 실제로 기록됐고(살균된 형태로) 검증됨
