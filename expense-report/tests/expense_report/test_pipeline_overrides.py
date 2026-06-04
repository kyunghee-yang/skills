"""run_pipeline 재실행 시 Rule8 수기 입력 보존 통합 커버리지(멱등 재실행)."""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "src"))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from expense_report.main import run_pipeline
from expense_report.config import (
    SHEET2_DATA_START_ROW, SHEET2_COL_ACCOUNT, SHEET2_COL_USAGE,
    SHEET2_COL_EXPENSE, SHEET2_COL_COMPANION,
)
from tests.fixtures import build_sample_xls  # type: ignore


def test_manual_overrides_preserved_on_rerun(tmp_path, sample_template):
    sub = tmp_path / "202603"
    sub.mkdir()
    build_sample_xls(str(sub / "간편서비스_승인내역.xls"))

    # 1차 실행 → Rule8(수기) 항목이 하나 이상 있어야 한다
    r1 = run_pipeline(str(sub), notion_data=None, template_path=sample_template)
    assert r1["manual_items"], "합성 데이터에 Rule8 수기 항목이 있어야 함"
    out = r1["created_file"]
    idx = r1["manual_items"][0]["row"] - 1          # all_transactions index
    row = SHEET2_DATA_START_ROW + idx               # sheet2 행

    # 사용자가 수기로 입력한 상황을 모사: 해당 Rule8 행에 수기값 기입
    wb = openpyxl.load_workbook(out)
    ws = wb["2.(기명카드)사용내역"]
    ws.cell(row, SHEET2_COL_ACCOUNT).value = "복리후생비[식비]"
    ws.cell(row, SHEET2_COL_USAGE).value = "수기 점심"
    ws.cell(row, SHEET2_COL_EXPENSE).value = 8000
    ws.cell(row, SHEET2_COL_COMPANION).value = "양경희"
    wb.save(out)

    # 2차 실행 → 기존 수기 입력이 보존되어야 한다(line 124 경로)
    run_pipeline(str(sub), notion_data=None, template_path=sample_template)
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2["2.(기명카드)사용내역"]
    assert ws2.cell(row, SHEET2_COL_ACCOUNT).value == "복리후생비[식비]"
    assert ws2.cell(row, SHEET2_COL_USAGE).value == "수기 점심"
    assert ws2.cell(row, SHEET2_COL_EXPENSE).value == 8000
    assert ws2.cell(row, SHEET2_COL_COMPANION).value == "양경희"
