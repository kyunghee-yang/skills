"""main._read_existing_overrides 커버리지(재실행 시 수기 입력 보존)."""

import openpyxl

from expense_report.main import _read_existing_overrides
from expense_report.config import (
    SHEET2_DATA_START_ROW, SHEET2_COL_ACCOUNT, SHEET2_COL_USAGE,
    SHEET2_COL_EXPENSE, SHEET2_COL_COMPANION,
)

SHEET2 = "2.(기명카드)사용내역"


def _make_output(path, rows):
    """rows: list of (account, usage, expense, companion) or None for blank account row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET2
    r = SHEET2_DATA_START_ROW
    for row in rows:
        if row is not None:
            account, usage, expense, companion = row
            ws.cell(r, SHEET2_COL_ACCOUNT).value = account
            ws.cell(r, SHEET2_COL_USAGE).value = usage
            ws.cell(r, SHEET2_COL_EXPENSE).value = expense
            ws.cell(r, SHEET2_COL_COMPANION).value = companion
        r += 1
    wb.save(path)


def test_returns_empty_when_file_missing(tmp_path):
    assert _read_existing_overrides(str(tmp_path / "nope.xlsx")) == {}


def test_extracts_manual_overrides(tmp_path):
    p = str(tmp_path / "out.xlsx")
    _make_output(p, [
        ("여비교통비[택시]", "야근 택시비", 8000, "양경희"),  # idx 0
        None,                                                  # idx 1: account 없음 → 제외
        ("복리후생비[식비]", "저녁", 12000, "양경희,확인필요"),  # idx 2
    ])
    ov = _read_existing_overrides(p)
    assert set(ov.keys()) == {0, 2}
    assert ov[0]["account"] == "여비교통비[택시]"
    assert ov[0]["usage"] == "야근 택시비"
    assert ov[0]["expense"] == 8000
    assert ov[2]["companion"] == "양경희,확인필요"
